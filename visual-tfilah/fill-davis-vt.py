"""
Fill the Davis VT frame with the Kabbalat Shabbat service content.

Input : working.pptx  (15 slides already ordered, all content slides cloned
                       from the template's authoring slide -- see build.sh)
Output: Davis-Kabbalat-Shabbat-VT-v1a.pptx, or argv[1]

Slide 1 (title) and slide 15 (closing) are Davis template slides, left untouched.
"""
import copy
import sys
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
q = lambda t: f"{{{A}}}{t}"

WHITE = "FFFFFF"
MUTED = "B8E3E6"  # light teal, for sidebar items that are not current
LINE_SP = 5760  # hundredths of a point; shared by Hebrew + transliteration

# Service order shown in the right-hand sidebar.
SIDEBAR = ["Bim Bam", "Video", "Last Word", "Candles", "Kiddush",
           "Hamotzi", "Oseh Shalom", "Priestly Blessing", "Rocks"]


# --------------------------------------------------------------------------
# XML helpers
# --------------------------------------------------------------------------
def walk(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:
            yield from walk(sh.shapes)


def by_name(slide, name):
    for sh in walk(slide.shapes):
        if sh.name == name:
            return sh
    return None


def drop(shape):
    """Remove a shape from its parent."""
    if shape is not None:
        shape._element.getparent().remove(shape._element)


def set_lines(shape, lines, line_spacing=None, color=None, size=None,
              italic=None, bold=None):
    """
    Replace a text box's content with `lines`, cloning the paragraph and run
    properties of its first paragraph so template formatting survives.
    """
    body = shape.text_frame._txBody
    paras = body.findall(q("p"))
    model = paras[0]
    model_pPr = model.find(q("pPr"))
    model_run = model.find(q("r"))
    model_rPr = model_run.find(q("rPr")) if model_run is not None else None

    for p in paras:
        body.remove(p)

    if not lines:
        lines = [""]

    for text in lines:
        p = body.makeelement(q("p"), {})
        if model_pPr is not None:
            pPr = copy.deepcopy(model_pPr)
            if line_spacing is not None:
                lnSpc = pPr.find(q("lnSpc"))
                if lnSpc is not None:
                    pPr.remove(lnSpc)
                lnSpc = pPr.makeelement(q("lnSpc"), {})
                pts = lnSpc.makeelement(q("spcPts"), {"val": str(line_spacing)})
                lnSpc.append(pts)
                pPr.insert(0, lnSpc)
            p.append(pPr)
        r = p.makeelement(q("r"), {})
        if model_rPr is not None:
            rPr = copy.deepcopy(model_rPr)
            # internal slide-jump links come from the template; drop them
            for h in rPr.findall(q("hlinkClick")):
                rPr.remove(h)
            if size is not None:
                rPr.set("sz", str(int(size * 100)))
            if italic is not None:
                rPr.set("i", "1" if italic else "0")
            if bold is not None:
                rPr.set("b", "1" if bold else "0")
            if color is not None:
                fill = rPr.find(q("solidFill"))
                if fill is not None:
                    rPr.remove(fill)
                fill = rPr.makeelement(q("solidFill"), {})
                clr = fill.makeelement(q("srgbClr"), {"val": color})
                fill.append(clr)
                # solidFill must precede the typeface elements
                idx = 0
                for child in rPr:
                    if child.tag in (q("latin"), q("ea"), q("cs"), q("sym"),
                                     q("rtl"), q("highlight"), q("uFill")):
                        break
                    idx += 1
                rPr.insert(idx, fill)
            r.append(rPr)
        t = r.makeelement(q("t"), {})
        t.text = text
        r.append(t)
        p.append(r)
        body.append(p)


def set_sidebar(slide, current):
    """Service-order list; the current item in white, the rest muted."""
    box = by_name(slide, "TextBox 13")
    if box is None:
        return
    # 20pt rather than the template's 24pt so "Priestly Blessing" stays on one line
    set_lines(box, SIDEBAR, color=MUTED, size=20, line_spacing=3000)
    body = box.text_frame._txBody
    for p, label in zip(body.findall(q("p")), SIDEBAR):
        if label != current:
            continue
        for r in p.findall(q("r")):
            rPr = r.find(q("rPr"))
            fill = rPr.find(q("solidFill"))
            if fill is not None:
                rPr.remove(fill)
            fill = rPr.makeelement(q("solidFill"), {})
            fill.append(fill.makeelement(q("srgbClr"), {"val": WHITE}))
            rPr.insert(0, fill)


# --------------------------------------------------------------------------
# Slide builders
# --------------------------------------------------------------------------
def add_dots(slide, index, total):
    """
    Progress dots for a prayer that runs across several slides. They sit in the
    header bar, in the clear span between the English title (ends at 6.86") and
    the Hebrew title (starts at 13.07"), so they never collide with either.
    """
    box = slide.shapes.add_textbox(Inches(8.0), Inches(0.07),
                                   Inches(4.0), Inches(0.61))
    tf = box.text_frame
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = " ".join("●" if i == index else "○"
                        for i in range(1, total + 1))
    run.font.size = Pt(22)
    run.font.name = "Calibri (MS)"
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def build(slide, *, en_header, he_header, translit=None, hebrew=None,
          translation=None, sidebar=None, notes=None, wide_text=None,
          dots=None):
    set_lines(by_name(slide, "TextBox 15"), [en_header])
    set_lines(by_name(slide, "TextBox 12"), [he_header])

    tlit = by_name(slide, "TextBox 10")
    heb = by_name(slide, "TextBox 11")

    if wide_text is not None:
        # Non-liturgy slide: one wide column across the content field.
        tlit.left, tlit.width = Inches(0.45), Inches(15.3)
        set_lines(tlit, wide_text, line_spacing=6400)
        drop(heb)
    else:
        # The template column is narrow enough that several transliteration
        # lines wrapped, which broke the line-for-line match with the Hebrew.
        # Widen it into the empty gutter; the Hebrew column starts at 9.48".
        tlit.width = Inches(8.5)
        set_lines(tlit, translit or [], line_spacing=LINE_SP)
        set_lines(heb, hebrew or [], line_spacing=LINE_SP)

    group = by_name(slide, "Group 16")
    if translation:
        set_lines(by_name(slide, "TextBox 18"), [translation])
    else:
        drop(group)

    drop(by_name(slide, "TextBox 14"))  # siddur page-number key, unused here
    set_sidebar(slide, sidebar)
    if dots:
        add_dots(slide, *dots)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes


DRAFT = ("  DRAFT PROMPTS - replace with the prompts from the printed outline.")

pres = Presentation("working.pptx")
s = pres.slides

# --- 2. Intro: Bim Bam (Micah) --------------------------------------------
build(s[1],
      en_header="Bim Bam",
      he_header="בִּים בַּם",
      hebrew=["בִּים בַּם, בִּים בִּים בִּים בַּם,",
              "בִּים בִּים בִּים בִּים בִּים בַּם.",
              "שַׁבָּת שָׁלוֹם!",
              "שַׁבָּת, שַׁבָּת, שַׁבָּת,",
              "שַׁבָּת שָׁלוֹם!"],
      translit=["Bim bam, bim bim bim bam,",
                "bim bim bim bim bim bam.",
                "Shabbat shalom!",
                "Shabbat, Shabbat, Shabbat,",
                "Shabbat shalom!"],
      translation="A peaceful Shabbat to you.",
      sidebar="Bim Bam",
      notes="Intro - Bim Bam. Leader: Micah. Opening niggun welcoming Shabbat. "
            "Music: Nurit Hirsh.")

# --- 3. Video (Cristy) ----------------------------------------------------
build(s[2],
      en_header="Video",
      he_header="",
      wide_text=[""],
      sidebar="Video",
      notes="Video. Leader: Cristy. Content area left open - place the video "
            "here and set it to play full screen.")

# --- 4. Last Word, round 1 (Micah facilitates) ----------------------------
build(s[3],
      en_header="Last Word",
      he_header="הַמִּלָּה הָאַחֲרוֹנָה",
      wide_text=["Something I am excited about this year is...",
                 "",
                 "As I look at the year ahead, I am hoping for..."],
      sidebar="Last Word",
      notes="Last Word sharing, round 1. Micah facilitates." + DRAFT)

# --- 5. Candles (Amy lights) ----------------------------------------------
build(s[4],
      en_header="Candle Blessing",
      he_header="הַדְלָקַת נֵרוֹת",
      hebrew=["בָּרוּךְ אַתָּה יְיָ,",
              "אֱלֹהֵינוּ מֶלֶךְ הָעוֹלָם,",
              "אֲשֶׁר קִדְּשָׁנוּ בְּמִצְוֹתָיו",
              "וְצִוָּנוּ לְהַדְלִיק",
              "נֵר שֶׁל שַׁבָּת."],
      translit=["Baruch atah Adonai,",
                "Eloheinu Melech ha'olam,",
                "asher kid'shanu b'mitzvotav",
                "v'tzivanu l'hadlik",
                "ner shel Shabbat."],
      translation="Blessed are You, Adonai our God, Sovereign of all, who makes us "
                  "holy with mitzvot and calls us to kindle the lights of Shabbat.",
      sidebar="Candles",
      notes="Candle blessing. Amy lights.")

# --- 6. Last Word, round 2 ------------------------------------------------
build(s[5],
      en_header="Last Word",
      he_header="הַמִּלָּה הָאַחֲרוֹנָה",
      wide_text=["A moment when a colleague really showed up for me was...",
                 "",
                 "Something my team is celebrating this year is..."],
      sidebar="Last Word",
      notes="Last Word sharing, round 2. Micah facilitates." + DRAFT)

# --- 7-9. Kiddush, full text over three slides (Dan Medwin) ---------------
build(s[6],
      en_header="Kiddush",
      he_header="קִדּוּשׁ",
      hebrew=["בָּרוּךְ אַתָּה יְיָ,",
              "אֱלֹהֵינוּ מֶלֶךְ הָעוֹלָם,",
              "בּוֹרֵא פְּרִי הַגָּפֶן."],
      translit=["Baruch atah Adonai,",
                "Eloheinu Melech ha'olam,",
                "borei p'ri hagafen."],
      translation="Blessed are You, Adonai our God, Sovereign of all, "
                  "Creator of the fruit of the vine.",
      sidebar="Kiddush", dots=(1, 3),
      notes="Kiddush 1 of 3, blessing over wine. Leader: Dan Medwin.")

build(s[7],
      en_header="Kiddush",
      he_header="קִדּוּשׁ",
      hebrew=["בָּרוּךְ אַתָּה יְיָ,",
              "אֱלֹהֵינוּ מֶלֶךְ הָעוֹלָם,",
              "אֲשֶׁר קִדְּשָׁנוּ בְּמִצְוֹתָיו",
              "וְרָצָה בָנוּ, וְשַׁבַּת קָדְשׁוֹ",
              "בְּאַהֲבָה וּבְרָצוֹן הִנְחִילָנוּ,",
              "זִכָּרוֹן לְמַעֲשֵׂה בְרֵאשִׁית.",
              "כִּי הוּא יוֹם תְּחִלָּה",
              "לְמִקְרָאֵי קֹדֶשׁ,",
              "זֵכֶר לִיצִיאַת מִצְרָיִם."],
      translit=["Baruch atah Adonai,",
                "Eloheinu Melech ha'olam,",
                "asher kid'shanu b'mitzvotav",
                "v'ratzah vanu, v'Shabbat kodsho",
                "b'ahavah uv'ratzon hinchilanu,",
                "zikaron l'maaseih v'reishit.",
                "Ki hu yom t'chilah",
                "l'mikra'ei kodesh,",
                "zeicher litziat Mitzrayim."],
      # The translation box holds three lines (~73 characters each). A fourth
      # runs off the bottom of the slide, so this one is condensed to fit.
      translation="Blessed are You, Adonai our God, Sovereign of all, who makes us "
                  "holy with mitzvot and delights in us, giving us Your holy Shabbat "
                  "in love, a reminder of creation and of the going out from Egypt.",
      sidebar="Kiddush", dots=(2, 3),
      notes="Kiddush 2 of 3. Leader: Dan Medwin.")

build(s[8],
      en_header="Kiddush",
      he_header="קִדּוּשׁ",
      hebrew=["כִּי בָנוּ בָחַרְתָּ",
              "וְאוֹתָנוּ קִדַּשְׁתָּ",
              "מִכָּל הָעַמִּים,",
              "וְשַׁבַּת קָדְשְׁךָ בְּאַהֲבָה",
              "וּבְרָצוֹן הִנְחַלְתָּנוּ.",
              "",
              "בָּרוּךְ אַתָּה יְיָ,",
              "מְקַדֵּשׁ הַשַּׁבָּת."],
      translit=["Ki vanu vacharta",
                "v'otanu kidashta",
                "mikol ha'amim,",
                "v'Shabbat kodsh'cha b'ahavah",
                "uv'ratzon hinchaltanu.",
                "",
                "Baruch atah Adonai,",
                "m'kadeish haShabbat."],
      translation="You have chosen us and made us holy among all peoples, and in "
                  "love and favor have given us Your holy Shabbat as a heritage. "
                  "Blessed are You, Adonai, who makes Shabbat holy.",
      sidebar="Kiddush", dots=(3, 3),
      notes="Kiddush 3 of 3, ending with the chatimah. Leader: Dan Medwin.")

# --- 10. Last Word, round 3 ------------------------------------------------
build(s[9],
      en_header="Last Word",
      he_header="הַמִּלָּה הָאַחֲרוֹנָה",
      wide_text=["A risk I would like to take this year is...",
                 "",
                 "The blessing I want to carry into this year is..."],
      sidebar="Last Word",
      notes="Last Word sharing, round 3. Micah facilitates." + DRAFT)

# --- 11. Hamotzi (Micah) ---------------------------------------------------
build(s[10],
      en_header="Hamotzi",
      he_header="הַמּוֹצִיא",
      hebrew=["בָּרוּךְ אַתָּה יְיָ,",
              "אֱלֹהֵינוּ מֶלֶךְ הָעוֹלָם,",
              "הַמּוֹצִיא לֶחֶם מִן הָאָרֶץ."],
      translit=["Baruch atah Adonai,",
                "Eloheinu Melech ha'olam,",
                "hamotzi lechem min ha'aretz."],
      translation="Blessed are You, Adonai our God, Sovereign of all, "
                  "who brings forth bread from the earth.",
      sidebar="Hamotzi",
      notes="Hamotzi, blessing over bread. Leader: Micah.")

# --- 12. Oseh Shalom (Micah) ----------------------------------------------
build(s[11],
      en_header="Oseh Shalom",
      he_header="עֹשֶׂה שָׁלוֹם",
      hebrew=["עֹשֶׂה שָׁלוֹם בִּמְרוֹמָיו,",
              "הוּא יַעֲשֶׂה שָׁלוֹם עָלֵינוּ,",
              "וְעַל כָּל יִשְׂרָאֵל,",
              "וְעַל כָּל יוֹשְׁבֵי תֵבֵל,",
              "וְאִמְרוּ: אָמֵן."],
      translit=["Oseh shalom bimromav,",
                "hu ya'aseh shalom aleinu,",
                "v'al kol Yisrael,",
                "v'al kol yosh'vei teiveil,",
                "v'imru: Amen."],
      translation="May the One who makes peace in the high heavens make peace for us, "
                  "for all Israel, and for all who dwell on earth. And let us say: Amen.",
      sidebar="Oseh Shalom",
      notes="Oseh Shalom. Leader: Micah.")

# --- 13. Priestly Blessing (Micah, Javier, Michelle, Emilie) ---------------
build(s[12],
      en_header="Priestly Blessing",
      he_header="בִּרְכַּת כֹּהֲנִים",
      hebrew=["יְבָרֶכְךָ יְיָ",
              "וְיִשְׁמְרֶךָ.",
              "יָאֵר יְיָ פָּנָיו אֵלֶיךָ",
              "וִיחֻנֶּךָּ.",
              "יִשָּׂא יְיָ פָּנָיו אֵלֶיךָ",
              "וְיָשֵׂם לְךָ שָׁלוֹם."],
      translit=["Y'varech'cha Adonai",
                "v'yishm'recha.",
                "Ya'er Adonai panav eilecha",
                "vichuneka.",
                "Yisa Adonai panav eilecha",
                "v'yasem l'cha shalom."],
      translation="May Adonai bless you and protect you. May Adonai's face shine upon "
                  "you and be gracious to you. May Adonai's face be lifted toward you "
                  "and grant you peace.",
      sidebar="Priestly Blessing",
      notes="Priestly Blessing, Numbers 6:24-26. Leaders: Micah, Javier, Michelle, "
            "and Emilie, one verse each with Micah leading.")

# --- 14. Rocks (Micah) ----------------------------------------------------
build(s[13],
      en_header="Rocks",
      he_header="אֲבָנִים",
      wide_text=[""],
      sidebar="Rocks",
      notes="Rocks slide. Leader: Micah. Content area left open - drop in the "
            "rocks image or text here.")

out = sys.argv[1] if len(sys.argv) > 1 else "Davis-Kabbalat-Shabbat-VT-v1a.pptx"
pres.save(out)
print("saved", out)
